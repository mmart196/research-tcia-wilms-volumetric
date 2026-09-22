"""Regression tests for documented Word formatting defects."""
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "code"))
from render_draft import write_docx


class WordRenderingTests(unittest.TestCase):
    def test_references_are_separate_paragraphs_and_affiliations_are_superscript(self):
        md = "Author ^1^\n\n## References\n\n1. First reference.\n2. Second reference.\n"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            write_docx(md, path)
            document = Document(path)
        self.assertEqual([p.text for p in document.paragraphs],
                         ["Author 1", "References", "1. First reference.", "2. Second reference."])
        marker = [r for r in document.paragraphs[0].runs if r.text == "1"]
        self.assertTrue(marker[0].font.superscript)


if __name__ == "__main__":
    unittest.main()
